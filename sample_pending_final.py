import sys, os
sys.path.insert(0, r'c:\Users\DELL\Desktop\daily_push')
sys.path.insert(0, r'c:\Users\DELL\Desktop\daily_push\push_bot')
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from generate_report import PENDING_REMARK_MAP

# Get actual values
deliver_val = PENDING_REMARK_MAP['306']
check_val = PENDING_REMARK_MAP['210']
return_val = PENDING_REMARK_MAP['500']

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Pending'

headers = ['ZONE', 'CURRENT POST OFFICE', 'ORDER ID', 'NEXT_ACTION', 'REMARK']
sample = [
    ['Zone 3', 'BATP001', '3103741408', deliver_val, deliver_val],
    ['Zone 3', 'BATP001', '3103742231', check_val, check_val],
    ['Zone 3', 'BATP001', '3103750875', deliver_val, deliver_val],
    ['Zone 3', 'BATP001', '3103752942', check_val, check_val],
    ['Zone 3', 'BATP001', '3103754119', check_val, check_val],
    ['Zone 3', 'BATP001', '3103757986', return_val, return_val],
    ['Zone 3', 'BATP001', '3103758586', check_val, check_val],
]

hfill = PatternFill('solid', fgColor='1F4E78')
hfont = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
thin = Side(style='thin', color='BFBFBF')
bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal='center', vertical='center')
red_fill = PatternFill('solid', fgColor='FFEBEB')
orange_fill = PatternFill('solid', fgColor='FCE4D6')
dfont = Font(name='Segoe UI', size=10)
bfont = Font(name='Segoe UI', size=10, bold=True)

for ci, h in enumerate(headers, 1):
    c = ws.cell(1, ci, h)
    c.fill = hfill
    c.font = hfont
    c.alignment = ctr
    c.border = bdr

for ri, row in enumerate(sample, 2):
    remark = row[4]
    if 'Return' in remark:
        fill = orange_fill
    else:
        fill = red_fill  # all urgent for demo
    for ci, v in enumerate(row, 1):
        c = ws.cell(ri, ci, v)
        c.border = bdr
        c.alignment = ctr
        c.fill = fill
        c.font = bfont

for ci, w in enumerate([9, 20, 16, 25, 25], 1):
    ws.column_dimensions[get_column_letter(ci)].width = w

out = r'c:\Users\DELL\Desktop\daily_push\sample_pending_final.xlsx'
wb.save(out)
print(f'Saved: {out}')

import excel_to_image
buf = excel_to_image.excel_to_image(out)
img_path = r'c:\Users\DELL\Desktop\daily_push\sample_pending_final.png'
with open(img_path, 'wb') as f:
    f.write(buf.getvalue())
print(f'Image: {img_path}')
