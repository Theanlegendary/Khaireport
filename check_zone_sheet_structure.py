"""Check ZONE sheet structure to find correct cell ranges for reports"""
import openpyxl
import sys
sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('0.Master Daily Report - new - Sept_New.xlsx', data_only=False)
ws = wb['ZONE']

print("=== ZONE Sheet Structure Analysis ===\n")

# Check for zone summary table
print("Looking for report tables...")
for r in range(1, 100):
    for c in range(1, 20):
        val = ws.cell(r, c).value
        if val and 'bao cao' in str(val).lower().replace('á', 'a').replace('ả', 'a'):
            print(f"Found at {ws.cell(r, c).coordinate}: {val}")

print("\n=== Checking rows 1-70 for table structures ===")
for r in range(1, 71):
    row_data = []
    has_content = False
    for c in range(1, 15):
        val = ws.cell(r, c).value
        if val and str(val).strip():
            has_content = True
            row_data.append(f"{ws.cell(r, c).coordinate}={str(val)[:30]}")
    
    if has_content:
        print(f"Row {r:2d}: {' | '.join(row_data[:5])}")

print("\n=== Looking for Customer Report table ===")
for r in range(1, 100):
    for c in range(1, 20):
        val = ws.cell(r, c).value
        if val and 'khach hang' in str(val).lower().replace('á', 'a').replace('ả', 'a'):
            print(f"Found customer report marker at {ws.cell(r, c).coordinate}")
            # Show next 10 rows
            print(f"  Next 10 rows starting from row {r}:")
            for r2 in range(r, min(r+10, ws.max_row)):
                cells = []
                for c2 in range(1, 13):
                    v = ws.cell(r2, c2).value
                    if v:
                        cells.append(f"{ws.cell(r2, c2).coordinate}={str(v)[:20]}")
                if cells:
                    print(f"    Row {r2}: {cells[:5]}")
