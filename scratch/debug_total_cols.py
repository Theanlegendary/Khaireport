import openpyxl

fpath = r"C:\Users\DELL\Downloads\Telegram Desktop\Total_13.07_14H32.xlsx"
wb = openpyxl.load_workbook(fpath)
ws = wb.active

with open("c:/Users/DELL/Desktop/daily_push/scratch/total_cols_debug.txt", "w", encoding="utf-8") as out:
    for r in range(1, 40):
        row_vals = [str(ws.cell(r, c).value) if ws.cell(r, c).value is not None else "" for c in range(1, ws.max_column + 1)]
        out.write(f"Row {r:02d}: {', '.join(row_vals)}\n")
        
print("SUCCESS")
