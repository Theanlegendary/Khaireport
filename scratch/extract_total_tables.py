import openpyxl

fpath = r"C:\Users\DELL\Downloads\Telegram Desktop\Total_13.07_14H32.xlsx"
wb = openpyxl.load_workbook(fpath)
ws = wb.active

print("Workbook sheets:", wb.sheetnames)

# We want to scan the active sheet vertically, find the table headers, 
# and count rows for Pickup, Delivery, and Pending.
sections = {}
current_section = None
section_row_count = 0

for r in range(1, ws.max_row + 1):
    val_a = ws.cell(r, 1).value
    val_d = ws.cell(r, 4).value # Column 4 is usually ORDER ID
    val_c = ws.cell(r, 3).value # Column 3 is usually ORDER ID in Pending
    
    val_a_str = str(val_a).strip() if val_a is not None else ""
    
    # Check if this row is a section title
    if "PICKUP" in val_a_str.upper() or "អីវ៉ាន់ត្រូវយក" in val_a_str:
        if current_section:
            sections[current_section] = section_row_count
        current_section = "Pickup"
        section_row_count = 0
        continue
    elif "DELIVERY" in val_a_str.upper() or "អីវ៉ាន់ត្រូវដឹក" in val_a_str:
        if current_section:
            sections[current_section] = section_row_count
        current_section = "Delivery"
        section_row_count = 0
        continue
    elif "PENDING" in val_a_str.upper() or "អីវ៉ាន់កំពុងរង់ចាំ" in val_a_str:
        if current_section:
            sections[current_section] = section_row_count
        current_section = "Pending"
        section_row_count = 0
        continue
        
    # If in a section, count data rows
    if current_section:
        # Check if the row contains an order ID (which is a digit)
        # It could be in Column 3 or 4
        is_order = False
        for c in (3, 4):
            v = ws.cell(r, c).value
            if v is not None and str(v).strip().isdigit():
                is_order = True
                break
        if is_order:
            section_row_count += 1

if current_section:
    sections[current_section] = section_row_count

print("Extracted counts from Total_13.07_14H32.xlsx:")
print(sections)
