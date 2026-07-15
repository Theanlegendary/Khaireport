import openpyxl
import os

out_dir = "c:/Users/DELL/Desktop/daily_push/test_out"
files_with_time = [
    (os.path.join(out_dir, f), os.path.getmtime(os.path.join(out_dir, f)))
    for f in os.listdir(out_dir)
    if os.path.isfile(os.path.join(out_dir, f))
]
files_with_time.sort(key=lambda x: x[1], reverse=True)
latest_files = [x[0] for x in files_with_time]

delivery_file = next((f for f in latest_files if "BANP001" in f and "Delivery" in f), None)

if delivery_file:
    print(f"Opening latest: {delivery_file}")
    wb = openpyxl.load_workbook(delivery_file)
    ws = wb.active
    
    # Iterate through rows and print cell color fills
    for r in range(1, ws.max_row + 1):
        cell_val = ws.cell(r, 4).value # ORDER ID
        fill_color = ws.cell(r, 4).fill.start_color.rgb
        print(f"Row {r} | Fill: {fill_color}")
else:
    print("Delivery file not found.")
