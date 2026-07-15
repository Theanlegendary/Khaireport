import openpyxl
import collections

file_morning = r"c:\Users\DELL\Desktop\daily_push - Copy\test_out\Report_All_13_07_2026.xlsx"
file_afternoon = r"C:\Users\DELL\Downloads\Telegram Desktop\Total_13.07_14H32.xlsx"

def clean_zone(z):
    if not z:
        return "Unknown"
    z_str = str(z).strip().title()
    if "Zone" in z_str:
        return z_str
    return f"Zone {z_str}"

# 1. Parse Morning
def get_morning_counts():
    wb = openpyxl.load_workbook(file_morning)
    counts = collections.defaultdict(int)
    for sname in wb.sheetnames:
        ws = wb[sname]
        key = "Pending"
        if "pickup" in sname.lower():
            key = "Pickup"
        elif "delivery" in sname.lower() or "deliver" in sname.lower():
            key = "Delivery"
            
        # Find order_id col
        order_col_idx = 4
        for c in range(1, ws.max_column + 1):
            cell_val = ws.cell(2, c).value
            if cell_val and any(h in str(cell_val).lower() for h in ('order id', 'លេខបុង')):
                order_col_idx = c
                break
                
        for r in range(4, ws.max_row + 1):
            order_val = ws.cell(r, order_col_idx).value
            if order_val is not None and str(order_val).strip().isdigit():
                zone_val = ws.cell(r, 1).value
                zone_name = clean_zone(zone_val)
                counts[(key, zone_name)] += 1
    return counts

# 2. Parse Afternoon
def get_afternoon_counts():
    wb = openpyxl.load_workbook(file_afternoon)
    ws = wb.active
    counts = collections.defaultdict(int)
    current_section = None
    r = 1
    while r <= ws.max_row:
        val_a = ws.cell(r, 1).value
        val_a_str = str(val_a).strip() if val_a is not None else ""
        
        if "PICKUP" in val_a_str.upper() or "អីវ៉ាន់ត្រូវយក" in val_a_str:
            current_section = "Pickup"
            r += 1
            continue
        elif "DELIVERY" in val_a_str.upper() or "អីវ៉ាន់ត្រូវដឹក" in val_a_str:
            current_section = "Delivery"
            r += 1
            continue
        elif "PENDING" in val_a_str.upper() or "អីវ៉ាន់កំពុងរង់ចាំ" in val_a_str:
            current_section = "Pending"
            r += 1
            continue
            
        col_4_val = ws.cell(r, 4).value
        if col_4_val and any(h in str(col_4_val) for h in ('ORDER ID', 'លេខបុង')):
            r += 2
            continue
            
        if current_section:
            order_id = ws.cell(r, 4).value
            if order_id is not None and str(order_id).strip().isdigit():
                zone_val = ws.cell(r, 1).value
                zone_name = clean_zone(zone_val)
                counts[(current_section, zone_name)] += 1
        r += 1
    return counts

morning_counts = get_morning_counts()
afternoon_counts = get_afternoon_counts()

zones = ["Zone 1", "Zone 2", "Zone 3", "Zone 4", "Zone 5"]
categories = ["Pickup", "Delivery", "Pending"]

def format_clearance_table(m_dict, a_dict, cat_name):
    out = []
    out.append(f"### 📍 {cat_name} Table\n")
    out.append(f"| Zone | 9:00 AM | 2:00 PM | Bills Cleared | % Cleared |\n")
    out.append(f"| :--- | :---: | :---: | :---: | :---: |\n")
    
    tot_m = 0
    tot_a = 0
    for z in zones:
        m = m_dict.get((cat_name, z), 0)
        a = a_dict.get((cat_name, z), 0)
        cleared = m - a
        pct = (cleared / m * 100) if m > 0 else 0.0
        
        # format positive/negative
        cleared_str = f"+{cleared}" if cleared > 0 else f"{cleared}"
        # If cleared is positive, it means bills went down (cleared)
        # If cleared is negative, it means bills went up (accumulated)
        # We can format it exactly as a raw number or with + / -
        pct_str = f"{pct:.2f}%"
        
        out.append(f"| {z} | {m} | {a} | **{cleared}** | **{pct_str}** |\n")
        tot_m += m
        tot_a += a
        
    tot_cleared = tot_m - tot_a
    tot_pct = (tot_cleared / tot_m * 100) if tot_m > 0 else 0.0
    out.append(f"| **TOTAL ALL ZONE** | **{tot_m}** | **{tot_a}** | **{tot_cleared}** | **{tot_pct:.2f}%** |\n\n")
    return "".join(out)

# Total overall zones (sum of all three categories)
def format_grand_total_table(m_dict, a_dict):
    out = []
    out.append(f"### 📍 Grand Total Table (All Categories Combined)\n")
    out.append(f"| Zone | 9:00 AM | 2:00 PM | Bills Cleared | % Cleared |\n")
    out.append(f"| :--- | :---: | :---: | :---: | :---: |\n")
    
    tot_m = 0
    tot_a = 0
    for z in zones:
        m = sum(m_dict.get((cat, z), 0) for cat in categories)
        a = sum(a_dict.get((cat, z), 0) for cat in categories)
        cleared = m - a
        pct = (cleared / m * 100) if m > 0 else 0.0
        
        out.append(f"| {z} | {m} | {a} | **{cleared}** | **{pct:.2f}%** |\n")
        tot_m += m
        tot_a += a
        
    tot_cleared = tot_m - tot_a
    tot_pct = (tot_cleared / tot_m * 100) if tot_m > 0 else 0.0
    out.append(f"| **TOTAL ALL ZONE** | **{tot_m}** | **{tot_a}** | **{tot_cleared}** | **{tot_pct:.2f}%** |\n\n")
    return "".join(out)

with open("c:/Users/DELL/Desktop/daily_push/scratch/zone_clearance_report.md", "w", encoding="utf-8") as f:
    f.write("# Zone-wise Clearance Report (9:00 AM vs 2:00 PM)\n\n")
    f.write("This report tracks the **Bills Cleared** and **% Cleared** by Zone for each category, using the exact formulas from your example:\n")
    f.write("- `Bills Cleared = 9:00 AM - 2:00 PM` (Positive means bills decreased/cleared, negative means bills increased)\n")
    f.write("- `% Cleared = (Bills Cleared / 9:00 AM) * 100` (Positive means clearance rate, negative means accumulation rate)\n\n")
    
    for cat in categories:
        f.write(format_clearance_table(morning_counts, afternoon_counts, cat))
        
    f.write(format_grand_total_table(morning_counts, afternoon_counts))

print("SUCCESS")
