import openpyxl

fpath = r"c:\Users\DELL\Desktop\daily_push - Copy\test_out\Report_All_13_07_2026.xlsx"
wb = openpyxl.load_workbook(fpath)

with open("c:/Users/DELL/Desktop/daily_push/scratch/morning_cols_debug.txt", "w", encoding="utf-8") as out:
    for sname in wb.sheetnames:
        out.write(f"\n==================== Sheet: {sname} ====================\n")
        ws = wb[sname]
        for r in range(1, 15):
            row_vals = [str(ws.cell(r, c).value) if ws.cell(r, c).value is not None else "" for c in range(1, ws.max_column + 1)]
            out.write(f"Row {r:02d}: {', '.join(row_vals)}\n")
            
print("SUCCESS")
