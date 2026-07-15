import openpyxl
import collections

file_morning = r"c:\Users\DELL\Desktop\daily_push - Copy\test_out\Report_All_13_07_2026.xlsx"
file_afternoon = r"C:\Users\DELL\Downloads\Telegram Desktop\Total_13.07_14H32.xlsx"

def clean_zone(z):
    if not z:
        return "Unknown"
    z_str = str(z).strip().title()
    if "Zone" in z_str:
        return z_str
    return f"Zone {z_str}"

def get_date_columns_map(ws):
    date_cols = {}
    current_month = None
    for c in range(1, ws.max_column + 1):
        m_val = ws.cell(2, c).value
        if m_val is not None:
            m_str = str(m_val).strip()
            if m_str in ("June", "July"):
                current_month = m_str
        d_val = ws.cell(3, c).value
        if d_val is not None and str(d_val).strip().isdigit() and current_month:
            date_cols[c] = (current_month, int(str(d_val).strip()))
    return date_cols

def is_urgent(month, day):
    if month == "June":
        return True
    if month == "July" and day <= 11:
        return True
    return False

# 1. Parse Morning
def parse_morning():
    wb = openpyxl.load_workbook(file_morning)
    total_counts = collections.defaultdict(int)  # {(cat, zone): count}
    urgent_counts = collections.defaultdict(int) # {(cat, zone): count}
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        cat = "Pending"
        if "pickup" in sname.lower():
            cat = "Pickup"
        elif "delivery" in sname.lower() or "deliver" in sname.lower():
            cat = "Delivery"
            
        date_cols = get_date_columns_map(ws)
        
        # Find order_id column header
        order_col_idx = 4
        for c in range(1, ws.max_column + 1):
            cell_val = ws.cell(2, c).value
            if cell_val and any(h in str(cell_val).lower() for h in ('order id', 'លេខបុង')):
                order_col_idx = c
                break
                
        for r in range(4, ws.max_row + 1):
            order_id = ws.cell(r, order_col_idx).value
            if order_id is not None and str(order_id).strip().isdigit():
                zone_val = ws.cell(r, 1).value
                zone_name = clean_zone(zone_val)
                
                # Total counts
                total_counts[(cat, zone_name)] += 1
                
                # Urgent counts
                created_date = None
                for c_idx, (m, d) in date_cols.items():
                    cell_val = ws.cell(r, c_idx).value
                    if cell_val == 1 or str(cell_val).strip() == '1':
                        created_date = (m, d)
                        break
                if created_date and is_urgent(created_date[0], created_date[1]):
                    urgent_counts[(cat, zone_name)] += 1
                    
    return total_counts, urgent_counts

# 2. Parse Afternoon
def parse_afternoon():
    wb = openpyxl.load_workbook(file_afternoon)
    ws = wb.active
    total_counts = collections.defaultdict(int)
    urgent_counts = collections.defaultdict(int)
    
    current_section = None
    date_cols = {}
    
    r = 1
    while r <= ws.max_row:
        val_a = ws.cell(r, 1).value
        val_a_str = str(val_a).strip() if val_a is not None else ""
        
        # Section headers
        if "PICKUP" in val_a_str.upper() or "អីវ៉ាន់ត្រូវយក" in val_a_str:
            current_section = "Pickup"
            date_cols = {}
            r += 1
            continue
        elif "DELIVERY" in val_a_str.upper() or "អីវ៉ាន់ត្រូវដឹក" in val_a_str:
            current_section = "Delivery"
            date_cols = {}
            r += 1
            continue
        elif "PENDING" in val_a_str.upper() or "អីវ៉ាន់កំពុងរង់ចាំ" in val_a_str:
            current_section = "Pending"
            date_cols = {}
            r += 1
            continue
            
        # Check if header row
        col_4_val = ws.cell(r, 4).value
        if col_4_val and any(h in str(col_4_val) for h in ('ORDER ID', 'លេខបុង')):
            r_months = r
            r_days = r + 1
            current_month = None
            for c in range(1, ws.max_column + 1):
                m_val = ws.cell(r_months, c).value
                if m_val is not None:
                    m_str = str(m_val).strip()
                    if m_str in ("June", "July"):
                        current_month = m_str
                d_val = ws.cell(r_days, c).value
                if d_val is not None and str(d_val).strip().isdigit() and current_month:
                    date_cols[c] = (current_month, int(str(d_val).strip()))
            r = r_days + 1
            continue
            
        if current_section:
            order_id = ws.cell(r, 4).value
            if order_id is not None and str(order_id).strip().isdigit():
                zone_val = ws.cell(r, 1).value
                zone_name = clean_zone(zone_val)
                
                # Total counts
                total_counts[(current_section, zone_name)] += 1
                
                # Urgent counts
                created_date = None
                for c_idx, (m, d) in date_cols.items():
                    cell_val = ws.cell(r, c_idx).value
                    if cell_val == 1 or str(cell_val).strip() == '1':
                        created_date = (m, d)
                        break
                if created_date and is_urgent(created_date[0], created_date[1]):
                    urgent_counts[(current_section, zone_name)] += 1
                    
        r += 1
    return total_counts, urgent_counts

morning_tot, morning_urg = parse_morning()
afternoon_tot, afternoon_urg = parse_afternoon()

zones = ["Zone 1", "Zone 2", "Zone 3", "Zone 4", "Zone 5"]
cats = ["Pickup", "Delivery", "Pending"]

def format_table(m_dict, a_dict, cat_name, title):
    out = []
    out.append(f"### 📍 {title} - {cat_name}\n")
    out.append(f"| Zone | 9:00 AM | 2:00 PM | Change (+/-) | % Change |\n")
    out.append(f"| :--- | :---: | :---: | :---: | :---: |\n")
    
    tot_m = 0
    tot_a = 0
    for z in zones:
        m = m_dict.get((cat_name, z), 0)
        a = a_dict.get((cat_name, z), 0)
        diff = a - m
        pct = (diff / m * 100) if m > 0 else (100.0 if a > 0 else 0.0)
        
        diff_str = f"+{diff}" if diff > 0 else f"{diff}"
        pct_str = f"+{pct:.2f}%" if diff > 0 else f"{pct:.2f}%"
        
        out.append(f"| {z} | {m} | {a} | **{diff_str}** | **{pct_str}** |\n")
        tot_m += m
        tot_a += a
        
    tot_diff = tot_a - tot_m
    tot_pct = (tot_diff / tot_m * 100) if tot_m > 0 else 0.0
    tot_diff_str = f"+{tot_diff}" if tot_diff > 0 else f"{tot_diff}"
    tot_pct_str = f"+{tot_pct:.2f}%" if tot_diff > 0 else f"{tot_pct:.2f}%"
    
    out.append(f"| **TOTAL ALL ZONE** | **{tot_m}** | **{tot_a}** | **{tot_diff_str}** | **{tot_pct_str}** |\n\n")
    return "".join(out)

# Build MD file
with open("c:/Users/DELL/Desktop/daily_push/scratch/zone_comparison_report.md", "w", encoding="utf-8") as f:
    f.write("# Zone-wise Report Comparison (9:00 AM vs 2:00 PM)\n\n")
    f.write("This report breaks down the morning and afternoon counts by **Zone** for **Pickup**, **Delivery**, and **Pending** orders.\n\n")
    
    f.write("## 1. TOTAL VOLUME TABLES\n")
    f.write("These tables show the **Total Volume** of orders in the system. Note that counts generally increased due to new orders imported today:\n\n")
    for cat in cats:
        f.write(format_table(morning_tot, afternoon_tot, cat, "Total Volume"))
        
    f.write("---\n\n")
    
    f.write("## 2. URGENT (OVERDUE >= 2 DAYS) TABLES\n")
    f.write("These tables show the **Urgent (Overdue)** orders. In some categories like **Pending**, you can see how many overdue bills were successfully cleared:\n\n")
    for cat in cats:
        f.write(format_table(morning_urg, afternoon_urg, cat, "Urgent Volume"))

print("SUCCESS")
