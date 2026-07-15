import openpyxl
import os

folder = r"C:\Users\DELL\Downloads\Telegram Desktop"
files = ["Total_13.07_14H32.xlsx", "Total_ZONE1_13.07_14H45.xlsx"]

with open("c:/Users/DELL/Desktop/daily_push/scratch/real_afternoon_counts.txt", "w", encoding="utf-8") as out:
    for f in files:
        fpath = os.path.join(folder, f)
        if not os.path.exists(fpath):
            out.write(f"\nFile {f} DOES NOT EXIST\n")
            continue
        try:
            wb = openpyxl.load_workbook(fpath)
            out.write(f"\nFile: {f}\n")
            for sname in wb.sheetnames:
                ws = wb[sname]
                
                # Find Order ID column
                header_row = None
                order_col_idx = None
                for r in range(1, 10):
                    for c in range(1, ws.max_column + 1):
                        val = ws.cell(r, c).value
                        if val and any(h in str(val) for h in ('ORDER ID', 'លេខបុង')):
                            header_row = r
                            order_col_idx = c
                            break
                    if header_row:
                        break
                
                if not header_row or not order_col_idx:
                    out.write(f"  Sheet: {repr(sname)} | No header found\n")
                    continue
                    
                cnt = 0
                for r in range(header_row + 1, ws.max_row + 1):
                    val = ws.cell(r, order_col_idx).value
                    if val is not None and str(val).strip().isdigit():
                        cnt += 1
                out.write(f"  Sheet: {repr(sname)} | Count: {cnt}\n")
        except Exception as e:
            out.write(f"Error reading {f}: {e}\n")

print("SUCCESS")
