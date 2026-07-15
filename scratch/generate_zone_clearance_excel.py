import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
import os

# Data
zones = ["Zone 1", "Zone 2", "Zone 3", "Zone 4", "Zone 5"]
categories = ["Pickup", "Delivery", "Pending"]

morning_data = {
    'Pickup': {"Zone 1": 12, "Zone 2": 14, "Zone 3": 11, "Zone 4": 0, "Zone 5": 12},
    'Delivery': {"Zone 1": 68, "Zone 2": 34, "Zone 3": 24, "Zone 4": 16, "Zone 5": 7},
    'Pending': {"Zone 1": 178, "Zone 2": 161, "Zone 3": 271, "Zone 4": 81, "Zone 5": 57}
}

afternoon_data = {
    'Pickup': {"Zone 1": 47, "Zone 2": 36, "Zone 3": 43, "Zone 4": 11, "Zone 5": 17},
    'Delivery': {"Zone 1": 140, "Zone 2": 63, "Zone 3": 32, "Zone 4": 15, "Zone 5": 42},
    'Pending': {"Zone 1": 431, "Zone 2": 286, "Zone 3": 142, "Zone 4": 152, "Zone 5": 159}
}

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Clearance Summary"
ws.views.sheetView[0].showGridLines = True

# Styles
font_title = Font(name="SF Pro Text", size=16, bold=True, color="1F4E78")
font_section = Font(name="SF Pro Text", size=12, bold=True, color="2F5597")
font_header = Font(name="SF Pro Text", size=10, bold=True, color="FFFFFF")
font_data = Font(name="SF Pro Text", size=10)
font_total = Font(name="SF Pro Text", size=10, bold=True)

fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
fill_total = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
fill_green = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
fill_red = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")

thin_side = Side(border_style="thin", color="D9D9D9")
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_total = Border(left=thin_side, right=thin_side, top=Side(border_style="thin", color="000000"), bottom=Side(border_style="double", color="000000"))

align_center = Alignment(horizontal="center", vertical="center")
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

# Title Block
ws.cell(1, 1, "ZONE-WISE CLEARANCE REPORT (9:00 AM vs 2:00 PM)").font = font_title
ws.cell(2, 1, "Report Date: July 13th, 2026").font = Font(name="SF Pro Text", size=10, italic=True)

current_row = 4

def write_table(ws, cat_name, m_dict, a_dict, start_row):
    # Table Header
    ws.cell(start_row, 1, f"📍 {cat_name.upper()} CLEARANCE TABLE").font = font_section
    
    headers = ["Zone", "9:00 AM", "2:00 PM", "Bills Cleared", "% Cleared"]
    for c, h in enumerate(headers, 1):
        cell = ws.cell(start_row + 1, c, h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell
        
    tot_m = 0
    tot_a = 0
    
    row_idx = start_row + 2
    for z in zones:
        m = m_dict.get(z, 0)
        a = a_dict.get(z, 0)
        cleared = m - a
        pct = (cleared / m) if m > 0 else 0.0
        
        ws.cell(row_idx, 1, z).font = font_data
        ws.cell(row_idx, 1).alignment = align_left
        ws.cell(row_idx, 1).border = border_cell
        
        ws.cell(row_idx, 2, m).font = font_data
        ws.cell(row_idx, 2).alignment = align_right
        ws.cell(row_idx, 2).border = border_cell
        ws.cell(row_idx, 2).number_format = "#,##0"
        
        ws.cell(row_idx, 3, a).font = font_data
        ws.cell(row_idx, 3).alignment = align_right
        ws.cell(row_idx, 3).border = border_cell
        ws.cell(row_idx, 3).number_format = "#,##0"
        
        # Bills Cleared Formula
        ws.cell(row_idx, 4, f"=B{row_idx}-C{row_idx}").font = font_data
        ws.cell(row_idx, 4).alignment = align_right
        ws.cell(row_idx, 4).border = border_cell
        ws.cell(row_idx, 4).number_format = "+#,##0;-#,##0;0"
        
        # % Cleared Formula
        ws.cell(row_idx, 5, f"=IF(B{row_idx}>0, D{row_idx}/B{row_idx}, 0)").font = font_data
        ws.cell(row_idx, 5).alignment = align_right
        ws.cell(row_idx, 5).border = border_cell
        ws.cell(row_idx, 5).number_format = "0.00%"
        
        # Color Highlights
        if cleared > 0:
            ws.cell(row_idx, 4).fill = fill_green
            ws.cell(row_idx, 5).fill = fill_green
        elif cleared < 0:
            ws.cell(row_idx, 4).fill = fill_red
            ws.cell(row_idx, 5).fill = fill_red
            
        tot_m += m
        tot_a += a
        row_idx += 1
        
    # Total Row
    ws.cell(row_idx, 1, "TOTAL ALL ZONE").font = font_total
    ws.cell(row_idx, 1).alignment = align_left
    ws.cell(row_idx, 1).border = border_total
    ws.cell(row_idx, 1).fill = fill_total
    
    ws.cell(row_idx, 2, f"=SUM(B{start_row+2}:B{row_idx-1})").font = font_total
    ws.cell(row_idx, 2).alignment = align_right
    ws.cell(row_idx, 2).border = border_total
    ws.cell(row_idx, 2).fill = fill_total
    ws.cell(row_idx, 2).number_format = "#,##0"
    
    ws.cell(row_idx, 3, f"=SUM(C{start_row+2}:C{row_idx-1})").font = font_total
    ws.cell(row_idx, 3).alignment = align_right
    ws.cell(row_idx, 3).border = border_total
    ws.cell(row_idx, 3).fill = fill_total
    ws.cell(row_idx, 3).number_format = "#,##0"
    
    ws.cell(row_idx, 4, f"=B{row_idx}-C{row_idx}").font = font_total
    ws.cell(row_idx, 4).alignment = align_right
    ws.cell(row_idx, 4).border = border_total
    ws.cell(row_idx, 4).fill = fill_total
    ws.cell(row_idx, 4).number_format = "+#,##0;-#,##0;0"
    
    ws.cell(row_idx, 5, f"=IF(B{row_idx}>0, D{row_idx}/B{row_idx}, 0)").font = font_total
    ws.cell(row_idx, 5).alignment = align_right
    ws.cell(row_idx, 5).border = border_total
    ws.cell(row_idx, 5).fill = fill_total
    ws.cell(row_idx, 5).number_format = "0.00%"
    
    # Conditional format for Total Cleared
    tot_cleared = tot_m - tot_a
    if tot_cleared > 0:
        ws.cell(row_idx, 4).fill = fill_green
        ws.cell(row_idx, 5).fill = fill_green
    elif tot_cleared < 0:
        ws.cell(row_idx, 4).fill = fill_red
        ws.cell(row_idx, 5).fill = fill_red
        
    return row_idx + 3

# Write Individual Tables
current_row = write_table(ws, "Pickup", morning_data['Pickup'], afternoon_data['Pickup'], current_row)
current_row = write_table(ws, "Delivery", morning_data['Delivery'], afternoon_data['Delivery'], current_row)
current_row = write_table(ws, "Pending", morning_data['Pending'], afternoon_data['Pending'], current_row)

# Write Grand Total Table
ws.cell(current_row, 1, "📍 GRAND TOTAL CLEARANCE TABLE").font = font_section
headers = ["Zone", "9:00 AM", "2:00 PM", "Bills Cleared", "% Cleared"]
for c, h in enumerate(headers, 1):
    cell = ws.cell(current_row + 1, c, h)
    cell.font = font_header
    cell.fill = fill_header
    cell.alignment = align_center
    cell.border = border_cell
    
row_idx = current_row + 2
for z in zones:
    m = sum(morning_data[cat][z] for cat in categories)
    a = sum(afternoon_data[cat][z] for cat in categories)
    cleared = m - a
    
    ws.cell(row_idx, 1, z).font = font_data
    ws.cell(row_idx, 1).alignment = align_left
    ws.cell(row_idx, 1).border = border_cell
    
    ws.cell(row_idx, 2, m).font = font_data
    ws.cell(row_idx, 2).alignment = align_right
    ws.cell(row_idx, 2).border = border_cell
    ws.cell(row_idx, 2).number_format = "#,##0"
    
    ws.cell(row_idx, 3, a).font = font_data
    ws.cell(row_idx, 3).alignment = align_right
    ws.cell(row_idx, 3).border = border_cell
    ws.cell(row_idx, 3).number_format = "#,##0"
    
    ws.cell(row_idx, 4, f"=B{row_idx}-C{row_idx}").font = font_data
    ws.cell(row_idx, 4).alignment = align_right
    ws.cell(row_idx, 4).border = border_cell
    ws.cell(row_idx, 4).number_format = "+#,##0;-#,##0;0"
    
    ws.cell(row_idx, 5, f"=IF(B{row_idx}>0, D{row_idx}/B{row_idx}, 0)").font = font_data
    ws.cell(row_idx, 5).alignment = align_right
    ws.cell(row_idx, 5).border = border_cell
    ws.cell(row_idx, 5).number_format = "0.00%"
    
    if cleared > 0:
        ws.cell(row_idx, 4).fill = fill_green
        ws.cell(row_idx, 5).fill = fill_green
    elif cleared < 0:
        ws.cell(row_idx, 4).fill = fill_red
        ws.cell(row_idx, 5).fill = fill_red
    row_idx += 1

# Grand Total Row
ws.cell(row_idx, 1, "TOTAL ALL ZONE").font = font_total
ws.cell(row_idx, 1).alignment = align_left
ws.cell(row_idx, 1).border = border_total
ws.cell(row_idx, 1).fill = fill_total

ws.cell(row_idx, 2, f"=SUM(B{current_row+2}:B{row_idx-1})").font = font_total
ws.cell(row_idx, 2).alignment = align_right
ws.cell(row_idx, 2).border = border_total
ws.cell(row_idx, 2).fill = fill_total
ws.cell(row_idx, 2).number_format = "#,##0"

ws.cell(row_idx, 3, f"=SUM(C{current_row+2}:C{row_idx-1})").font = font_total
ws.cell(row_idx, 3).alignment = align_right
ws.cell(row_idx, 3).border = border_total
ws.cell(row_idx, 3).fill = fill_total
ws.cell(row_idx, 3).number_format = "#,##0"

ws.cell(row_idx, 4, f"=B{row_idx}-C{row_idx}").font = font_total
ws.cell(row_idx, 4).alignment = align_right
ws.cell(row_idx, 4).border = border_total
ws.cell(row_idx, 4).fill = fill_total
ws.cell(row_idx, 4).number_format = "+#,##0;-#,##0;0"

ws.cell(row_idx, 5, f"=IF(B{row_idx}>0, D{row_idx}/B{row_idx}, 0)").font = font_total
ws.cell(row_idx, 5).alignment = align_right
ws.cell(row_idx, 5).border = border_total
ws.cell(row_idx, 5).fill = fill_total
ws.cell(row_idx, 5).number_format = "0.00%"

# Adjust column widths
for col in ws.columns:
    max_len = 0
    for cell in col:
        if cell.value:
            max_len = max(max_len, len(str(cell.value)))
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 3, 14)

# Save paths
path_desktop = r"c:\Users\DELL\Desktop\daily_push\Zone_Clearance_Report_13_07_2026.xlsx"
path_downloads = r"C:\Users\DELL\Downloads\Telegram Desktop\Zone_Clearance_Report_13_07_2026.xlsx"

wb.save(path_desktop)
try:
    wb.save(path_downloads)
    print(f"Excel saved to Desktop and Downloads successfully!")
except Exception as e:
    print(f"Excel saved to Desktop. Downloads failed: {e}")
