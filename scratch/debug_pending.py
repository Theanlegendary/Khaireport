import openpyxl

fpath = r"C:\Users\DELL\Downloads\Telegram Desktop\Report_All_13_07_2026 (7).xlsx"
wb = openpyxl.load_workbook(fpath)

with open("c:/Users/DELL/Desktop/daily_push/scratch/pending_sheet_debug.txt", "w", encoding="utf-8") as out:
    for sname in wb.sheetnames:
        if "Transit" in sname or "Needed" in sname:
            out.write(f"\n=========================================\n")
            out.write(f"Sheet Name: {sname}\n")
            ws = wb[sname]
            out.write(f"Max Row: {ws.max_row} | Max Col: {ws.max_column}\n")
            for r in range(1, min(20, ws.max_row + 1)):
                row_vals = [str(ws.cell(r, c).value) if ws.cell(r, c).value is not None else "" for c in range(1, ws.max_column + 1)]
                out.write(f"Row {r:02d}: {', '.join(row_vals)}\n")
print("SUCCESS")
