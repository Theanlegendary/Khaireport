import openpyxl

file_morning = r"c:\Users\DELL\Desktop\daily_push - Copy\test_out\Report_All_13_07_2026.xlsx"
file_afternoon = r"C:\Users\DELL\Downloads\Telegram Desktop\Report_All_13_07_2026 (7).xlsx"

def get_sheet_count(fpath, sheet_name):
    wb = openpyxl.load_workbook(fpath)
    if sheet_name not in wb.sheetnames:
        return 0
    ws = wb[sheet_name]
    
    # Find header row and column for Order ID
    header_row = None
    order_col_idx = None
    for r in range(1, 10):
        for c in range(1, ws.max_column + 1):
            val = ws.cell(r, c).value
            if val and any(h in str(val) for h in ('ORDER ID', 'លេខបុង')):
                header_row = r
                order_col_idx = c
                break
        if header_row:
            break
            
    if not header_row or not order_col_idx:
        return 0
        
    cnt = 0
    for r in range(header_row + 1, ws.max_row + 1):
        val = ws.cell(r, order_col_idx).value
        if val is not None:
            val_str = str(val).strip()
            if val_str and val_str.isdigit():
                cnt += 1
    return cnt

# Morning counts
morning = {
    'Pickup': get_sheet_count(file_morning, 'Pickup'),
    'Delivery': get_sheet_count(file_morning, 'Delivery'),
    'Pending': get_sheet_count(file_morning, 'Pending')
}

# Afternoon counts (Note: Pending was split into 'កំពុងលើផ្លូវ (Transit)' and 'ដល់ប៉ុស្តិ៍ (Action Needed)')
afternoon = {
    'Pickup': get_sheet_count(file_afternoon, 'Pickup'),
    'Delivery': get_sheet_count(file_afternoon, 'Delivery'),
    'Pending': get_sheet_count(file_afternoon, 'កំពុងលើផ្លូវ (Transit)') + get_sheet_count(file_afternoon, 'ដល់ប៉ុស្តិ៍ (Action Needed)')
}

with open("c:/Users/DELL/Desktop/daily_push/scratch/comparison_summary.txt", "w", encoding="utf-8") as f:
    f.write("=== REPORT COMPARISON: 9:00 AM vs 2:00 PM (July 13th, 2026) ===\n\n")
    f.write(f"Morning File: {file_morning}\n")
    f.write(f"Afternoon File: {file_afternoon}\n\n")
    
    f.write(f"| Category | Morning (9 AM) | Afternoon (2 PM) | Change | Change % |\n")
    f.write(f"| :--- | :---: | :---: | :---: | :---: |\n")
    for k in ['Pickup', 'Delivery', 'Pending']:
        m = morning[k]
        a = afternoon[k]
        diff = a - m
        pct = (diff / m * 100) if m > 0 else 0.0
        f.write(f"| {k} | {m} | {a} | {diff:+} | {pct:+.2f}% |\n")
        
    m_tot = sum(morning.values())
    a_tot = sum(afternoon.values())
    tot_diff = a_tot - m_tot
    tot_pct = (tot_diff / m_tot * 100) if m_tot > 0 else 0.0
    f.write(f"| **Total** | **{m_tot}** | **{a_tot}** | **{tot_diff:+}** | **{tot_pct:+.2f}%** |\n")

print("SUCCESS - Comparison saved to scratch/comparison_summary.txt")
