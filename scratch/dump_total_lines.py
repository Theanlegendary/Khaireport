import openpyxl

fpath = r"C:\Users\DELL\Downloads\Telegram Desktop\Total_13.07_14H32.xlsx"
wb = openpyxl.load_workbook(fpath)
ws = wb.active

with open("c:/Users/DELL/Desktop/daily_push/scratch/total_lines_debug.txt", "w", encoding="utf-8") as out:
    for r in range(1, ws.max_row + 1):
        val_a = ws.cell(r, 1).value
        if val_a is not None:
            val_a_str = str(val_a).strip()
            # If it looks like a section header (contains "BILL CHECK" or Khmer titles)
            if any(x in val_a_str for x in ("BILL CHECK", "សរុប", "Grand Total", "អីវ៉ាន់")):
                out.write(f"Row {r:04d}: {val_a_str}\n")
                
print("SUCCESS")
