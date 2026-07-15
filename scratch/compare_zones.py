import openpyxl
import collections

file_morning = r"c:\Users\DELL\Desktop\daily_push - Copy\test_out\Report_All_13_07_2026.xlsx"
file_afternoon = r"C:\Users\DELL\Downloads\Telegram Desktop\Total_13.07_14H32.xlsx"

# Standardizing Zone Names
def clean_zone(z):
    if not z:
        return "Unknown"
    z_str = str(z).strip().title()
    # Normalize e.g. "Zone 1" -> "Zone 1"
    if "Zone" in z_str:
        return z_str
    return f"Zone {z_str}"

# 1. Parse Morning
def get_morning_counts():
    wb = openpyxl.load_workbook(file_morning)
    # counts[(category, zone)] = count
    counts = collections.defaultdict(int)
    
    for sname in wb.sheetnames:
        ws = wb[sname]
        key = "Pending"
        if "pickup" in sname.lower():
            key = "Pickup"
        elif "delivery" in sname.lower() or "deliver" in sname.lower():
            key = "Delivery"
            
        # Find headers
        header_row = 2
        order_col_idx = 4
        for c in range(1, ws.max_column + 1):
            cell_val = ws.cell(2, c).value
            if cell_val and any(h in str(cell_val).lower() for h in ('order id', 'លេខបុង')):
                order_col_idx = c
                break
                
        for r in range(4, ws.max_row + 1):
            order_val = ws.cell(r, order_col_idx).value
            if order_val is not None and str(order_val).strip().isdigit():
                zone_val = ws.cell(r, 1).value
                zone_name = clean_zone(zone_val)
                counts[(key, zone_name)] += 1
    return counts

# 2. Parse Afternoon (Stacked Total)
def get_afternoon_counts():
    wb = openpyxl.load_workbook(file_afternoon)
    ws = wb.active
    counts = collections.defaultdict(int)
    
    current_section = None
    r = 1
    while r <= ws.max_row:
        val_a = ws.cell(r, 1).value
        val_a_str = str(val_a).strip() if val_a is not None else ""
        
        # Section changes
        if "PICKUP" in val_a_str.upper() or "អីវ៉ាន់ត្រូវយក" in val_a_str:
            current_section = "Pickup"
            r += 1
            continue
        elif "DELIVERY" in val_a_str.upper() or "អីវ៉ាន់ត្រូវដឹក" in val_a_str:
            current_section = "Delivery"
            r += 1
            continue
        elif "PENDING" in val_a_str.upper() or "អីវ៉ាន់កំពុងរង់ចាំ" in val_a_str:
            current_section = "Pending"
            r += 1
            continue
            
        # Check if header row
        col_4_val = ws.cell(r, 4).value
        if col_4_val and any(h in str(col_4_val) for h in ('ORDER ID', 'លេខបុង')):
            r += 2 # skip header and day rows
            continue
            
        if current_section:
            # We assume order ID is in column 4 (if not header / total)
            order_id = ws.cell(r, 4).value
            if order_id is not None and str(order_id).strip().isdigit():
                zone_val = ws.cell(r, 1).value
                zone_name = clean_zone(zone_val)
                counts[(current_section, zone_name)] += 1
        r += 1
    return counts

morning_counts = get_morning_counts()
afternoon_counts = get_afternoon_counts()

print("Morning Counts:")
for k, v in sorted(morning_counts.items()):
    print(f"  {k}: {v}")
    
print("\nAfternoon Counts:")
for k, v in sorted(afternoon_counts.items()):
    print(f"  {k}: {v}")
