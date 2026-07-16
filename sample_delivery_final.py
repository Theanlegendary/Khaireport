import sys, os
sys.path.insert(0, r'c:\Users\DELL\Desktop\daily_push\push_bot')
import openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from generate_report import DELIVERY_ACTION_MAP

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Delivery'

headers = ['ZONE', 'POST OFFICE HANDLE', 'CURRENT POST OFFICE', 'ORDER ID', 'RECEIVER', 'ACTION', 'NEXT_STEP']
# Sample data with different statuses
rows = [
    ['Zone 3', 'BATP001', 'BATP001', '3103741408', 'Sokha', '400'],
    ['Zone 3', 'BATP001', 'BATP001', '3103742231', 'Vanna', '401'],
    ['Zone 3', 'BATP001', 'BATP001', '3103750875', 'Pheap', '402'],
    ['Zone 3', 'BATP001', 'BATP001', '3103752942', 'Dara', '420'],
    ['Zone 3', 'BATP001', 'BATP001', '3103754119', 'Maly', '430'],
    ['Zone 3', 'BATP001', 'BATP001', '3103757986', 'Rith', '460'],
    ['Zone 3', 'BATP001', 'BATP001', '3103758586', 'Thy', '470'],
    ['Zone 3', 'BATP001', 'BATP001', '3103758708', 'Kosal', '471'],
    ['Zone 3', 'BATP001', 'BATP001', '3103758719', 'Chhay', '472'],
    ['Zone 3', 'BATP001', 'BATP001', '3103758853', 'Nary', '480'],
]

hfill = PatternFill('solid', fgColor='1F4E78')
hfont = Font(name='Segoe UI', size=10, bold=True, color='FFFFFF')
thin = Side(style='thin', color='BFBFBF')
bdr = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal='center', vertical='center')
red_fill = PatternFill('solid', fgColor='FFEBEB')
orange_fill = PatternFill('solid', fgColor='FCE4D6')
dfont = Font(name='Segoe UI', size=10)
bfont = Font(name='Segoe UI', size=10, bold=True)

for ci, h in enumerate(headers, 1):
    c = ws.cell(1, ci, h)
    c.fill = hfill
    c.font = hfont
    c.alignment = ctr
    c.border = bdr

for ri, row in enumerate(rows, 2):
    zone, handle, po, oid, recv, sc = row
    action, next_step = DELIVERY_ACTION_MAP.get(sc, ('', ''))
    values = [zone, handle, po, oid, recv, action, next_step]
    for ci, v in enumerate(values, 1):
        c = ws.cell(ri, ci, v)
        c.border = bdr
        c.alignment = ctr
        c.font = dfont
        c.fill = red_fill  # all urgent for demo

for ci, w in enumerate([9, 18, 20, 16, 14, 16, 20], 1):
    ws.column_dimensions[get_column_letter(ci)].width = w

out = r'c:\Users\DELL\Desktop\daily_push\sample_delivery_final.xlsx'
wb.save(out)
print(f'Saved: {out}')

import excel_to_image
buf = excel_to_image.excel_to_image(out)
img_path = r'c:\Users\DELL\Desktop\daily_push\sample_delivery_final.png'
with open(img_path, 'wb') as f:
    f.write(buf.getvalue())
print(f'Image: {img_path}')
