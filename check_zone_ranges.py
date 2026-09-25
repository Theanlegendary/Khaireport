import openpyxl
import sys

# Load workbook
wb = openpyxl.load_workbook('0.Master Daily Report - new - Sept_New.xlsx', data_only=True)
ws = wb['ZONE']

print("=" * 80)
print("ZONE SHEET INSPECTION")
print("=" * 80)

print(f"\nSheet dimensions: {ws.dimensions}")
print(f"Max row: {ws.max_row}, Max column: {ws.max_column}")

print("\n" + "=" * 80)
print("CURRENT ZONE SUMMARY RANGE (B6:M14):")
print("=" * 80)
for i in range(6, 15):
    row_data = []
    for j in range(2, 14):  # B to M
        cell = ws.cell(row=i, column=j)
        row_data.append(f"{cell.coordinate}:{cell.value}")
    print(f"Row {i}: {' | '.join(row_data)}")

print("\n" + "=" * 80)
print("CURRENT CUSTOMER REPORT RANGE (B56:L64):")
print("=" * 80)
for i in range(56, 65):
    row_data = []
    for j in range(2, 13):  # B to L
        cell = ws.cell(row=i, column=j)
        row_data.append(f"{cell.coordinate}:{cell.value}")
    print(f"Row {i}: {' | '.join(row_data)}")

print("\n" + "=" * 80)
print("SCANNING FOR TABLE HEADERS (rows 1-70):")
print("=" * 80)
for i in range(1, 71):
    for j in range(1, 20):
        cell = ws.cell(row=i, column=j)
        if cell.value and isinstance(cell.value, str):
            val_lower = str(cell.value).lower()
            if any(keyword in val_lower for keyword in ['zone', 'summary', 'báo cáo', 'khách hàng', 'phát sinh', 'sản lượng']):
                print(f"{cell.coordinate}: {cell.value}")

wb.close()
