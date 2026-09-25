import openpyxl

wb = openpyxl.load_workbook('0.Master Daily Report - new - Sept_New.xlsx', read_only=True)

print('Checking ZONE and 5 Zone sheets structure:\n')

# Check ZONE sheet
if 'ZONE' in wb.sheetnames:
    ws = wb['ZONE']
    print('=== ZONE Sheet ===')
    print(f'Max Row: {ws.max_row}, Max Col: {ws.max_column}')
    print(f'Sample cells:')
    for r in range(1, min(15, ws.max_row + 1)):
        for c in range(1, min(15, ws.max_column + 1)):
            val = ws.cell(r, c).value
            if val and str(val).strip():
                print(f'  {ws.cell(r, c).coordinate} = {val}')
    print()

# Check 5 Zone sheet  
if '5 Zone' in wb.sheetnames:
    ws = wb['5 Zone']
    print('=== 5 Zone Sheet ===')
    print(f'Max Row: {ws.max_row}, Max Col: {ws.max_column}')
    print(f'Sample cells (first 50):')
    count = 0
    for r in range(1, min(20, ws.max_row + 1)):
        for c in range(1, min(15, ws.max_column + 1)):
            val = ws.cell(r, c).value
            if val and str(val).strip():
                print(f'  {ws.cell(r, c).coordinate} = {val}')
                count += 1
                if count >= 50:
                    break
        if count >= 50:
            break
