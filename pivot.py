"""
pivot.py
Doc file Excel chi tiet don (export-detail) va dung bang pivot
"PENDING BILL CHECK" giong mau:
  - Loc: CURRENT STATUS thuoc danh sach pending_status_codes
         (+ loai test neu exclude_test)
  - Hang (rows): ZONE > CURRENT POST OFFICE > ORDER ID
  - Cot (columns): MONTH / DAY theo CREATED DATE
  - Gia tri: Count of ORDER ID
  - Co dong/cot Grand Total
"""

from datetime import datetime
from collections import defaultdict, OrderedDict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ---- vi tri cot trong file nguon (0-based) ----
COL_CREATED_DATE = 1     # CREATED DATE  (dd/mm/yyyy HH:MM:SS)
COL_ORDER_ID = 2         # ORDER ID
COL_CURRENT_PO = 15      # CURRENT POST OFFICE
COL_CURRENT_STATUS = 23  # CURRENT STATUS  ("110 - Chua tiep nhan")
COL_SENDER = 3
COL_RECEIVER = 4


def _status_code(value):
    """Lay ma so dau chuoi trang thai: '110 - Chua tiep nhan' -> '110'."""
    if value is None:
        return ""
    s = str(value).strip()
    if " - " in s:
        s = s.split(" - ", 1)[0]
    return s.split()[0].strip() if s else ""


def _parse_day(value):
    """Tra ve (month, day) tu CREATED DATE. Ho tro dd/mm/yyyy ..."""
    if value is None:
        return None, None
    if isinstance(value, datetime):
        return value.month, value.day
    s = str(value).strip()
    date_part = s.split(" ")[0]
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%m/%d/%Y", "%d-%m-%Y"):
        try:
            dt = datetime.strptime(date_part, fmt)
            return dt.month, dt.day
        except ValueError:
            continue
    return None, None


def _zone_for(po_code, zone_cfg):
    if not po_code:
        return zone_cfg.get("default_zone", "Khac")
    po = str(po_code).strip()
    by_po = zone_cfg.get("by_post_office", {})
    if po in by_po:
        return by_po[po]
    by_prefix = zone_cfg.get("by_prefix", {})
    for prefix, zone in by_prefix.items():
        if po.upper().startswith(prefix.upper()):
            return zone
    return zone_cfg.get("default_zone", "Khac")


def _is_test_row(row, test_keywords):
    blob = " ".join(str(row[c] or "") for c in (COL_SENDER, COL_RECEIVER)).lower()
    return any(k.lower() in blob for k in test_keywords)


def read_source(path):
    """Doc file Excel nguon -> list rows (tuple), bo qua header."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    return rows[1:]  # bo header


def build_pivot(rows, pivot_cfg, zone_cfg):
    """
    Tra ve cau truc pivot:
      records: dict[zone][po][order_id] = day (str 'DD')
      days: danh sach cac ngay (sorted)
      month: thang xuat hien
    """
    pending = set(str(c).strip() for c in pivot_cfg.get("pending_status_codes", []))
    exclude_test = pivot_cfg.get("exclude_test", False)
    test_keywords = pivot_cfg.get("test_keywords", ["test"])

    tree = defaultdict(lambda: defaultdict(dict))  # zone -> po -> {order_id: day}
    days_seen = set()
    months_seen = set()

    for row in rows:
        if not row or row[COL_ORDER_ID] in (None, ""):
            continue
        if pending and _status_code(row[COL_CURRENT_STATUS]) not in pending:
            continue
        if exclude_test and _is_test_row(row, test_keywords):
            continue

        month, day = _parse_day(row[COL_CREATED_DATE])
        if day is None:
            continue
        po = str(row[COL_CURRENT_PO] or "").strip() or "(trong)"
        zone = _zone_for(po, zone_cfg)
        order_id = str(row[COL_ORDER_ID]).strip()

        tree[zone][po][order_id] = f"{day:02d}"
        days_seen.add(day)
        months_seen.add(month)

    days = sorted(days_seen)
    month = sorted(months_seen)[0] if months_seen else None
    return tree, days, month


# ---- styling ----
_HDR_FILL = PatternFill("solid", fgColor="DCE6F1")
_ZONE_FILL = PatternFill("solid", fgColor="EAEAEA")
_THIN = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTER = Alignment(horizontal="center", vertical="center")
_LEFT = Alignment(horizontal="left", vertical="center")
_RED = Font(color="C00000", bold=True)


def export_pivot(tree, days, month, pivot_cfg, out_path):
    """Xuat pivot ra file Excel giong mau."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (pivot_cfg.get("title", "PENDING BILL CHECK") or "REPORT")[:31]

    day_cols = [f"{d:02d}" for d in days]
    # bo cuc cot: A=ZONE, B=CURRENT POST OFFICE, C=ORDER ID,
    # D..=cac ngay, cuoi=Grand Total
    first_day_col = 4
    n_days = len(day_cols)
    gt_col = first_day_col + n_days

    # --- bo loc o tren cung (giong mau) ---
    ws.cell(1, 1, "Phan loai").font = Font(bold=True)
    ws.cell(1, 2, pivot_cfg.get("classification", "Post office"))
    ws.cell(2, 1, "Is test").font = Font(bold=True)
    ws.cell(2, 2, "#N/A" if pivot_cfg.get("exclude_test") else "(tat ca)")
    ws.cell(3, 1, "is ton ket noi").font = Font(bold=True)
    ws.cell(3, 2, pivot_cfg.get("is_label", "Pending"))
    title = ws.cell(4, 1, f"{pivot_cfg.get('title','PENDING BILL CHECK')}  "
                          f"{datetime.now():%d.%m_%HH%M}")
    title.font = Font(bold=True, color="C00000", size=12)

    # --- header bang ---
    r_count = 5
    r_month = 6
    r_head = 7
    ws.cell(r_count, 1, "Count of ORDER ID").font = Font(bold=True)
    mcell = ws.cell(r_count, first_day_col, "MONTH"); mcell.font = Font(bold=True)
    dcell = ws.cell(r_count, first_day_col + 1, "DAY"); dcell.font = Font(bold=True)
    if month is not None:
        mc = ws.cell(r_month, first_day_col, f"{month:02d}")
        mc.font = Font(bold=True); mc.fill = _HDR_FILL; mc.alignment = _CENTER
    ws.cell(r_count, gt_col, "Grand Total").font = Font(bold=True)

    for label, col in (("ZONE", 1), ("CURRENT POST OFFICE", 2), ("ORDER ID", 3)):
        c = ws.cell(r_head, col, label)
        c.font = Font(bold=True); c.fill = _HDR_FILL; c.border = _BORDER
    for i, d in enumerate(day_cols):
        c = ws.cell(r_head, first_day_col + i, d)
        c.font = Font(bold=True); c.fill = _HDR_FILL; c.alignment = _CENTER; c.border = _BORDER
    gtc = ws.cell(r_head, gt_col, "Grand Total")
    gtc.font = Font(bold=True); gtc.fill = _HDR_FILL; gtc.alignment = _CENTER; gtc.border = _BORDER

    # --- body ---
    r = r_head + 1
    col_totals = defaultdict(int)
    grand_total = 0

    for zone in sorted(tree.keys()):
        zone_first_row = r
        zone_written = False
        for po in sorted(tree[zone].keys()):
            po_written = False
            for order_id in sorted(tree[zone][po].keys()):
                day = tree[zone][po][order_id]
                zc = ws.cell(r, 1, zone if not zone_written else None)
                pc = ws.cell(r, 2, po if not po_written else None)
                oc = ws.cell(r, 3, order_id)
                zc.fill = _ZONE_FILL
                for c in (zc, pc, oc):
                    c.border = _BORDER
                oc.alignment = _LEFT
                # danh dau ngay
                for i, dlabel in enumerate(day_cols):
                    cell = ws.cell(r, first_day_col + i)
                    cell.border = _BORDER
                    cell.alignment = _CENTER
                    if dlabel == day:
                        cell.value = 1
                        cell.font = _RED
                        col_totals[dlabel] += 1
                tot = ws.cell(r, gt_col, 1)
                tot.alignment = _CENTER; tot.border = _BORDER
                grand_total += 1
                zone_written = True
                po_written = True
                r += 1

    # --- dong Grand Total ---
    gt_row = r
    ws.cell(gt_row, 1, "Grand Total").font = Font(bold=True)
    ws.cell(gt_row, 1).fill = _HDR_FILL
    for i, dlabel in enumerate(day_cols):
        c = ws.cell(gt_row, first_day_col + i, col_totals.get(dlabel, 0))
        c.font = Font(bold=True); c.alignment = _CENTER; c.fill = _HDR_FILL; c.border = _BORDER
    c = ws.cell(gt_row, gt_col, grand_total)
    c.font = Font(bold=True); c.alignment = _CENTER; c.fill = _HDR_FILL; c.border = _BORDER

    # --- do rong cot ---
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 14
    for i in range(n_days):
        ws.column_dimensions[get_column_letter(first_day_col + i)].width = 7
    ws.column_dimensions[get_column_letter(gt_col)].width = 12
    ws.freeze_panes = ws.cell(r_head + 1, 4)

    wb.save(out_path)
    return out_path, grand_total


def run(source_path, out_path, config):
    rows = read_source(source_path)
    tree, days, month = build_pivot(rows, config["pivot"], config["zone_mapping"])
    return export_pivot(tree, days, month, config["pivot"], out_path)


def _merge_pivot_cfg(base_pivot, report):
    """Gop cau hinh pivot goc voi 1 report (ghi de status codes / title / label)."""
    cfg = dict(base_pivot or {})
    cfg["pending_status_codes"] = report.get("status_codes",
                                             cfg.get("pending_status_codes", []))
    cfg["title"] = report.get("title", cfg.get("title", "REPORT"))
    cfg["is_label"] = report.get("is_label", report.get("title", "Pending"))
    return cfg


def run_report(rows, out_path, config, report):
    """Dung 1 report tu rows da doc san. Tra ve (out_path, total)."""
    pivot_cfg = _merge_pivot_cfg(config.get("pivot", {}), report)
    tree, days, month = build_pivot(rows, pivot_cfg, config["zone_mapping"])
    return export_pivot(tree, days, month, pivot_cfg, out_path)


if __name__ == "__main__":
    import sys, json
    cfg = json.load(open(sys.argv[3])) if len(sys.argv) > 3 else json.load(
        open("config.json"))
    out, total = run(sys.argv[1], sys.argv[2], cfg)
    print(f"Da xuat: {out}  (tong {total} don pending)")
